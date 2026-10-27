"""Turn one learner request into a study pack.

The page calls this when the learner presses the button. A second press
for the same unchanged page reads the saved pack and does not call the model.
"""

import logging
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass, replace
from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import urlparse

from docx import Document
from docx.document import Document as DocumentType

from studymaster.agents.client import Completer, OpenAIModelClient
from studymaster.agents.coordinator import (
    SECTION_KINDS,
    AgentSection,
    HiddenCheck,
    run_study_pack,
)
from studymaster.agents.errors import AgentError, EmptySourceError
from studymaster.agents.grader import SCORE_LABELS
from studymaster.agents.quotes import COVER
from studymaster.agents.lab import write_lab
from studymaster.agents.lesson import write_lesson
from studymaster.agents.outline import write_outline
from studymaster.agents.project import write_project
from studymaster.agents.request import StudyRequest
from studymaster.agents.runbook import write_runbook
from studymaster.agents.study import write_study
from studymaster.config import Settings, get_settings
from studymaster.extraction.errors import ReaderError
from studymaster.extraction.fetch import Article, read_page
from studymaster.storage.database import (
    LOCAL_LEARNER_ID,
    STATUS_ERROR,
    STATUS_OK,
    Resource,
    StudyStore,
    content_hash,
)

logger = logging.getLogger(__name__)

LEVELS = ("beginner", "intermediate", "advanced")
PASTED_TEXT_URL = "pasted-text"
UNREADABLE_TEXT = "This page did not include readable article text."
SHORTENED_NOTICE = "This page was shortened to 60,000 characters."
PageReader = Callable[..., Article]


@dataclass(frozen=True)
class StudyOutcome:
    """What the home page may show, plus answers that stay off the page."""

    ok: bool
    message: str
    title: str = ""
    outline: str = ""
    sections: tuple[AgentSection, ...] = ()
    hidden_answers: tuple[str, ...] = ()
    checks: tuple[HiddenCheck, ...] = ()
    resource_id: int = 0
    content_hash: str = ""
    source_url: str = ""
    from_saved: bool = False
    level: str = ""
    minutes: int = 0
    goal: str = ""
    saved_on: str = ""


def display_blocks(outcome: StudyOutcome) -> list[tuple[str, str]]:
    """Return the text the page is allowed to show.

    Lesson answers are not included.
    """
    if not outcome.ok:
        return [("error", outcome.message)]
    blocks: list[tuple[str, str]] = []
    if outcome.outline:
        blocks.append(("outline", outcome.outline))
    for section in outcome.sections:
        if section.status != STATUS_OK:
            blocks.append(("error", section.error_message))
        elif section.markdown:
            blocks.append((section.kind, section.markdown))
    return blocks


def pack_score_line(outcome: StudyOutcome) -> str:
    """Count sentences left out and questions that can still be graded."""
    left_out = 0
    for kind, text in display_blocks(outcome):
        if kind == "error":
            continue
        left_out += sum(1 for line in text.splitlines() if COVER in line)
    return f"Left out {left_out} sentences. Questions to check: {len(outcome.checks)}."


def study_page(
    *,
    url: str,
    level: str,
    goal: str,
    minutes: int,
    settings: Settings,
    store: StudyStore,
    read: PageReader = read_page,
    complete: Completer | None = None,
    remembered_answers: tuple[str, ...] = (),
    remembered_hash: str = "",
    article_text: str = "",
    on_progress: Callable[[str], None] | None = None,
    on_preview: Callable[[str], None] | None = None,
    owner_id: str = LOCAL_LEARNER_ID,
) -> StudyOutcome:
    """Read one page, or the pasted lesson, and return its study pack."""
    cleaned_url = url.strip()
    cleaned_goal = goal.strip()
    cleaned_level = level.strip().lower()
    _report(on_progress, "Checking the link.")
    pasted = _collapse_lines(article_text)
    if pasted:
        problem = _learner_problem(cleaned_level, cleaned_goal, minutes)
        if problem:
            return StudyOutcome(ok=False, message=problem)
        if len(pasted) < settings.min_source_chars:
            return StudyOutcome(ok=False, message=UNREADABLE_TEXT)
        _report(on_progress, "Using the pasted lesson.")
        shortened = len(pasted) > settings.max_source_chars
        text = pasted[: settings.max_source_chars] if shortened else pasted
        article = Article(
            url=PASTED_TEXT_URL,
            title=_pasted_title(text),
            headings=[],
            text=text,
            shortened=shortened,
        )
        source_url = PASTED_TEXT_URL
    else:
        problem = _request_problem(cleaned_url, cleaned_level, cleaned_goal, minutes)
        if problem:
            return StudyOutcome(ok=False, message=problem)
        _report(on_progress, "Reading the page.")
        try:
            article = read(cleaned_url, settings=settings)
        except ReaderError as exc:
            store.save_resource(
                url=cleaned_url,
                title="",
                status=STATUS_ERROR,
                error_message=str(exc),
                level=cleaned_level,
                goal=cleaned_goal,
                minutes=minutes,
                owner_id=owner_id,
            )
            return StudyOutcome(ok=False, message=str(exc))
        source_url = cleaned_url

    digest = content_hash(article.text)
    saved = store.find_success(
        url=source_url,
        extracted_text=article.text,
        level=cleaned_level,
        goal=cleaned_goal,
        minutes=minutes,
        owner_id=owner_id,
    )
    if saved is not None:
        _report(on_progress, "Opening the saved pack.")
        stored = tuple(
            HiddenCheck(row.question, row.answer, row.excerpt)
            for row in store.list_lesson_checks(saved.resource.id)
        )
        if stored:
            answers = tuple(item.answer for item in stored)
        elif remembered_hash == digest:
            answers = remembered_answers
        else:
            answers = ()
        sections = tuple(
            AgentSection(
                artifact.kind,
                artifact.status,
                artifact.error_message,
                artifact.markdown,
                [],
            )
            for artifact in saved.artifacts
        )
        return StudyOutcome(
            ok=True,
            message="",
            title=saved.resource.title,
            sections=sections,
            hidden_answers=answers,
            checks=stored,
            resource_id=saved.resource.id,
            content_hash=digest,
            source_url=saved.resource.url,
            from_saved=True,
            level=saved.resource.level,
            minutes=saved.resource.minutes,
            goal=saved.resource.goal,
            saved_on=saved.resource.fetched_at,
        )

    request = StudyRequest(
        title=article.title,
        url=source_url,
        headings=list(article.headings),
        source_text=article.text,
        level=cleaned_level,
        goal=cleaned_goal,
        minutes=minutes,
    )
    try:
        pack = run_study_pack(
            request,
            complete=complete,
            settings=settings,
            on_progress=on_progress,
            on_preview=_section_preview(on_preview, shortened=article.shortened),
        )
    except AgentError as exc:
        store.save_resource(
            url=source_url,
            title=article.title,
            status=STATUS_ERROR,
            error_message=str(exc),
            extracted_text=article.text,
            level=cleaned_level,
            goal=cleaned_goal,
            minutes=minutes,
            owner_id=owner_id,
        )
        return StudyOutcome(
            ok=False,
            message=str(exc),
            title=article.title,
            content_hash=digest,
        )

    resource = store.save_resource(
        url=source_url,
        title=article.title,
        status=STATUS_OK,
        extracted_text=article.text,
        level=cleaned_level,
        goal=cleaned_goal,
        minutes=minutes,
        owner_id=owner_id,
    )
    shown = tuple(
        _with_outline(section, pack.outline, shortened=article.shortened)
        for section in pack.sections
    )
    for section in shown:
        store.save_artifact(
            resource_id=resource.id,
            kind=section.kind,
            status=section.status,
            error_message=section.error_message,
            markdown=section.markdown,
            model_name=settings.openai_model,
        )
    lesson = next((section for section in shown if section.kind == "lesson"), None)
    if lesson is not None and lesson.hidden_checks:
        store.replace_lesson_checks(
            resource.id,
            [
                (item.question, item.answer, item.excerpt)
                for item in lesson.hidden_checks
            ],
        )
    answers = next(
        (section.hidden_answers for section in pack.sections if section.kind == "lesson"),
        (),
    )
    checks = lesson.hidden_checks if lesson is not None else ()
    outcome = StudyOutcome(
        ok=True,
        message="",
        title=article.title,
        sections=shown,
        hidden_answers=answers,
        checks=checks,
        resource_id=resource.id,
        content_hash=digest,
        source_url=source_url,
        level=cleaned_level,
        minutes=minutes,
        goal=cleaned_goal,
        saved_on=resource.fetched_at,
    )
    logger.info("%s", pack_score_line(outcome))
    return outcome


def _request_problem(url: str, level: str, goal: str, minutes: int) -> str:
    if not url:
        return "Paste a public page link."
    return _learner_problem(level, goal, minutes)


def _learner_problem(level: str, goal: str, minutes: int) -> str:
    if level not in LEVELS:
        return "Choose beginner, intermediate, or advanced."
    if not goal:
        return "Say what you want to be able to do."
    if minutes < 1:
        return "Enter at least 1 minute."
    return ""


def _report(on_progress: Callable[[str], None] | None, message: str) -> None:
    if on_progress is not None:
        on_progress(message)


def _collapse_lines(text: str) -> str:
    """Join extra spaces and drop blank lines."""
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _pasted_title(text: str) -> str:
    """Use the first line, cut at 80 characters."""
    first = text.splitlines()[0].strip() if text.splitlines() else ""
    if not first:
        return "Pasted article"
    return first[:80]


def _section_preview(
    on_preview: Callable[[str], None] | None,
    *,
    shortened: bool,
) -> Callable[[str, str], None] | None:
    """Turn a finished outline or section into the text shown during the wait."""
    if on_preview is None:
        return None

    def reveal(kind: str, text: str) -> None:
        body = _live_outline(text, shortened=shortened) if kind == "outline" else text.strip()
        if body:
            on_preview(body)

    return reveal


def _live_outline(outline: str, *, shortened: bool) -> str:
    """Show the outline as soon as it is written, including a shortened-page line."""
    notice = f"{SHORTENED_NOTICE}\n\n" if shortened else ""
    body = outline.strip()
    if not body:
        return notice.strip()
    return f"{notice}## Outline\n\n{body}"


def _with_outline(
    section: AgentSection,
    outline: str,
    *,
    shortened: bool = False,
) -> AgentSection:
    """Keep the shared outline in the study section so a saved pack still shows it."""
    if section.kind != "study" or not outline.strip():
        return section
    notice = f"{SHORTENED_NOTICE}\n\n" if shortened else ""
    markdown = f"{notice}## Outline\n\n{outline.strip()}\n\n{section.markdown}"
    return AgentSection(
        section.kind,
        section.status,
        section.error_message,
        markdown,
        section.checks,
        section.hidden_answers,
        section.hidden_checks,
    )


@dataclass(frozen=True)
class WaitingQuestion:
    """One lesson question that does not yet have a correct score."""

    resource_id: int
    title: str
    question: str


def waiting_questions(
    store: StudyStore,
    *,
    owner_id: str = LOCAL_LEARNER_ID,
) -> list[WaitingQuestion]:
    """Return questions that still need a correct reply for one learner.

    A correct score takes the question off the list. This reads the saved
    checks and attempts only. It does not fetch a page or call the model.
    """
    waiting: list[WaitingQuestion] = []
    for resource in saved_studies(store, owner_id=owner_id):
        correct = {
            item.question
            for item in store.list_attempts(resource.id)
            if item.score == "correct"
        }
        for check in store.list_lesson_checks(resource.id):
            if check.question in correct:
                continue
            waiting.append(
                WaitingQuestion(
                    resource_id=resource.id,
                    title=resource.title,
                    question=check.question,
                )
            )
    return waiting


def focus_waiting_question(outcome: StudyOutcome, question: str) -> StudyOutcome:
    """Move the chosen question to the front of the open pack."""
    chosen = [item for item in outcome.checks if item.question == question]
    if not chosen:
        return outcome
    rest = tuple(item for item in outcome.checks if item.question != question)
    return replace(outcome, checks=tuple(chosen) + rest)


def saved_studies(
    store: StudyStore,
    *,
    owner_id: str = LOCAL_LEARNER_ID,
) -> list[Resource]:
    """Return one learner's successful saved packs, newest first."""
    return [
        resource
        for resource in store.list_resources(owner_id=owner_id)
        if resource.status == STATUS_OK
    ]


def matching_studies(
    studies: list[Resource],
    query: str,
    pieces: dict[int, str] | None = None,
    lengths: dict[int, str] | None = None,
    sites: dict[int, str] | None = None,
    days: dict[int, str] | None = None,
    levels: dict[int, str] | None = None,
    scores: dict[int, str] | None = None,
    checks: dict[int, str] | None = None,
    progress: dict[int, str] | None = None,
    questions: dict[int, str] | None = None,
    replies: dict[int, str] | None = None,
    remaining: dict[int, str] | None = None,
    empty: dict[int, str] | None = None,
    runbooks: dict[int, str] | None = None,
    projects: dict[int, str] | None = None,
) -> list[Resource]:
    """Keep studies whose title, goal, pieces, minutes, site, day, level, score, check, progress, question, reply, left-to-check, empty-lesson, no-procedure, or uncovered-project line contains the words.

    An empty box keeps the whole list. Capital letters and extra spaces are ignored.
    A number must match a whole number, so 1 leaves Checked 0 of 12 off the search.
    Typing correct leaves Last check: Incorrect off the search.
    Typing purpose keeps the pack whose latest question says purpose.
    Typing reply keeps the pack whose reply day says reply.
    Typing left to check keeps a pack that says how many questions are still waiting.
    Typing empty keeps a pack whose lesson came back empty.
    Typing procedure keeps a pack that says this page has no procedure.
    Typing covered keeps a pack whose project is not covered by the page.
    """
    needle = " ".join(query.casefold().split())
    if not needle:
        return list(studies)
    piece_lines = pieces or {}
    length_lines = lengths or {}
    site_lines = sites or {}
    day_lines = days or {}
    level_lines = levels or {}
    score_lines = scores or {}
    check_lines = checks or {}
    progress_lines = progress or {}
    question_lines = questions or {}
    reply_lines = replies or {}
    remaining_lines = remaining or {}
    empty_lines = empty or {}
    runbook_lines = runbooks or {}
    project_lines = projects or {}
    matched: list[Resource] = []
    for resource in studies:
        title = " ".join((resource.title or resource.url).casefold().split())
        goal = " ".join(resource.goal.casefold().split())
        piece_line = " ".join(piece_lines.get(resource.id, "").casefold().split())
        length_line = length_lines.get(resource.id, "").casefold()
        site_line = " ".join(site_lines.get(resource.id, "").casefold().split())
        day_line = day_lines.get(resource.id, "").casefold()
        level_line = " ".join(level_lines.get(resource.id, "").casefold().split())
        score_line = score_lines.get(resource.id, "").casefold()
        check_line = check_lines.get(resource.id, "").casefold()
        progress_line = progress_lines.get(resource.id, "").casefold()
        question_line = question_lines.get(resource.id, "").casefold()
        reply_line = reply_lines.get(resource.id, "").casefold()
        remaining_line = remaining_lines.get(resource.id, "").casefold()
        empty_line = empty_lines.get(resource.id, "").casefold()
        runbook_line = runbook_lines.get(resource.id, "").casefold()
        project_line = project_lines.get(resource.id, "").casefold()
        if (
            needle in title
            or needle in goal
            or (piece_line and needle in piece_line)
            or _numbered_line_match(length_line, needle)
            or (site_line and needle in site_line)
            or _numbered_line_match(day_line, needle)
            or (level_line and needle in level_line)
            or _score_line_match(score_line, needle)
            or _check_line_match(check_line, needle)
            or _score_line_match(progress_line, needle)
            or _check_line_match(question_line, needle)
            or _check_line_match(reply_line, needle)
            or _score_line_match(remaining_line, needle)
            or _check_line_match(empty_line, needle)
            or _check_line_match(runbook_line, needle)
            or _check_line_match(project_line, needle)
        ):
            matched.append(resource)
    return matched


def _check_line_match(line: str, needle: str) -> bool:
    """Return whether the check line contains the words as whole words.

    Typing partial matches Last check: Partial. Typing correct leaves
    Last check: Incorrect off the search.
    """
    cleaned = line
    for mark in ".:;!?":
        cleaned = cleaned.replace(mark, " ")
    cleaned = " ".join(cleaned.split())
    if not cleaned or not needle:
        return False
    return f" {needle} " in f" {cleaned} "


def _score_line_match(line: str, needle: str) -> bool:
    """Return whether the score line contains the words.

    Typing 5 matches Questions to check: 5. A colon or period does not glue that 5
    to the word beside it. Typing 5 leaves 77 and 12 off the search.
    """
    cleaned = line
    for mark in ".:;!?":
        cleaned = cleaned.replace(mark, " ")
    return _numbered_line_match(cleaned, needle)


def _numbered_line_match(line: str, needle: str) -> bool:
    """Return whether the line contains the words.

    Typing 5 matches 5 minutes. It leaves 15 minutes and October 25 off the search.
    A comma does not glue the day to the year.
    """
    cleaned = " ".join(line.replace(",", " ").split())
    if not cleaned or needle not in cleaned:
        return False
    if any(character.isdigit() for character in needle):
        return f" {needle} " in f" {cleaned} "
    return True


_MONTHS = (
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)


def read_from_line(url: str) -> str:
    """Return the site name for a public page, such as Read from learn.microsoft.com.

    A pasted lesson has no site, so the line stays blank.
    """
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return ""
    host = parsed.hostname.lower().rstrip(".")
    return f"Read from {host}"


def saved_minutes(minutes: int) -> str:
    """Return the study length, such as 15 minutes. Zero stays blank."""
    if minutes == 1:
        return "1 minute"
    if minutes > 1:
        return f"{minutes} minutes"
    return ""


def saved_level(level: str) -> str:
    """Return who the pack was written for, such as For a beginner.

    A missing level stays blank.
    """
    cleaned = level.strip().lower()
    if cleaned == "beginner":
        return "For a beginner"
    if cleaned == "intermediate":
        return "For an intermediate learner"
    if cleaned == "advanced":
        return "For an advanced learner"
    return ""


def saved_check(score: str) -> str:
    """Return the latest check, such as Last check: Partial.

    A pack with no check stays blank.
    """
    label = SCORE_LABELS.get(score.strip().lower(), "")
    if not label:
        return ""
    return f"Last check: {label}"


def saved_question(question: str) -> str:
    """Return the latest question, such as Last question: What does the service do?

    A pack with no check stays blank. The hint and the reply stay off this line.
    """
    cleaned = " ".join(question.split())
    if not cleaned:
        return ""
    return f"Last question: {cleaned}"


def saved_progress(answered: int, total: int) -> str:
    """Return how many questions have a reply, such as Checked 1 of 6.

    A pack with no questions stays blank. A second reply to the same question
    still counts as one question.
    """
    if total < 1:
        return ""
    count = min(max(answered, 0), total)
    return f"Checked {count} of {total}."


def left_to_check(answered: int, total: int) -> str:
    """Return how many questions are still waiting, such as Left to check: 5.

    A pack with no questions stays blank. A finished pack says Left to check: 0.
    A second reply to the same question still counts as one question.
    """
    if total < 1:
        return ""
    count = min(max(answered, 0), total)
    return f"Left to check: {total - count}."


def saved_pieces(kinds: Iterable[str]) -> str:
    """Return the saved sections, such as Study, lesson.

    The names follow the pack order. A missing section stays off the line.
    """
    present = set(kinds)
    names = [kind for kind in SECTION_KINDS if kind in present]
    if not names:
        return ""
    line = ", ".join(names)
    return line[0].upper() + line[1:]


def saved_day(fetched_at: str) -> str:
    """Return the saved day on this computer, such as Saved October 25, 2026."""
    try:
        parsed = datetime.fromisoformat(fetched_at)
    except ValueError:
        return ""
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    local = parsed.astimezone()
    return f"Saved {_MONTHS[local.month - 1]} {local.day}, {local.year}"


def reply_day(created_at: str) -> str:
    """Return the day of the latest reply, such as Reply day: October 25, 2026.

    A pack with no check stays blank. The hint and the reply stay off this line.
    """
    try:
        parsed = datetime.fromisoformat(created_at)
    except ValueError:
        return ""
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    local = parsed.astimezone()
    return f"Reply day: {_MONTHS[local.month - 1]} {local.day}, {local.year}."


def open_saved(
    store: StudyStore,
    resource_id: int,
    *,
    owner_id: str = LOCAL_LEARNER_ID,
) -> StudyOutcome:
    """Load one saved pack. This does not fetch the page or call the model."""
    resource = store.get_resource(resource_id)
    if resource.owner_id != owner_id:
        return StudyOutcome(ok=False, message="This study belongs to another learner.")
    if resource.status != STATUS_OK:
        return StudyOutcome(
            ok=False,
            message=resource.error_message or "This page was not saved.",
            title=resource.title,
            resource_id=resource.id,
            content_hash=resource.content_hash,
        )
    stored = tuple(
        HiddenCheck(row.question, row.answer, row.excerpt)
        for row in store.list_lesson_checks(resource.id)
    )
    sections = tuple(
        AgentSection(
            artifact.kind,
            artifact.status,
            artifact.error_message,
            artifact.markdown,
            [],
        )
        for artifact in store.list_artifacts(resource.id)
    )
    return StudyOutcome(
        ok=True,
        message="",
        title=resource.title,
        sections=sections,
        hidden_answers=tuple(item.answer for item in stored),
        checks=stored,
        resource_id=resource.id,
        content_hash=resource.content_hash,
        source_url=resource.url,
        from_saved=True,
        level=resource.level,
        minutes=resource.minutes,
        goal=resource.goal,
        saved_on=resource.fetched_at,
    )


def pack_markdown(outcome: StudyOutcome) -> str:
    """Return the downloadable pack. Lesson answers are omitted."""
    lines = [f"# {outcome.title}".rstrip(), ""]
    for _kind, text in display_blocks(outcome):
        lines.extend([text, ""])
    return "\n".join(lines).strip() + "\n"


_HEADING_LINE = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")
_NUMBERED_LINE = re.compile(r"^(\d+)\.\s+(.*\S)\s*$")
_BULLET_LINE = re.compile(r"^[-*]\s+(.*\S)\s*$")


def pack_docx(outcome: StudyOutcome) -> bytes:
    """Return the pack as a Word file. Lesson answers are omitted."""
    document = Document()
    title = outcome.title.strip()
    if title:
        document.add_heading(title, level=1)
    for _kind, text in display_blocks(outcome):
        _add_markdown_block(document, text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_markdown_block(document: DocumentType, text: str) -> None:
    """Add one section. Blank lines separate paragraphs."""
    for raw in text.splitlines():
        if not raw.strip():
            continue
        heading = _HEADING_LINE.match(raw.strip())
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_heading(heading.group(2).replace("**", ""), level=level)
            continue
        if raw[:1].isdigit():
            numbered = _NUMBERED_LINE.match(raw.strip())
            if numbered:
                paragraph = document.add_paragraph(style="List Number")
                _write_runs(paragraph, numbered.group(2))
                continue
        if raw.lstrip()[:1] in "-*":
            bullet = _BULLET_LINE.match(raw.strip())
            if bullet:
                paragraph = document.add_paragraph(style="List Bullet")
                _write_runs(paragraph, bullet.group(1))
                continue
        paragraph = document.add_paragraph()
        _write_runs(paragraph, raw.strip())


def _write_runs(paragraph: object, text: str) -> None:
    """Write **bold** words as bold text. The rest stays plain."""
    for index, part in enumerate(text.split("**")):
        if not part:
            continue
        run = paragraph.add_run(part)
        if index % 2 == 1:
            run.bold = True


def markdown_filename(title: str) -> str:
    """Return a safe file name for a downloaded pack."""
    cleaned = "".join(character if character.isalnum() else "-" for character in title.lower())
    cleaned = "-".join(part for part in cleaned.split("-") if part)
    return f"{cleaned or 'study'}.md"


def word_filename(title: str) -> str:
    """Return a safe Word file name for a downloaded pack."""
    return f"{markdown_filename(title).removesuffix('.md')}.docx"


def regenerate_section(
    *,
    store: StudyStore,
    resource_id: int,
    kind: str,
    settings: Settings | None = None,
    complete: Completer | None = None,
) -> StudyOutcome:
    """Rewrite one saved section. The page is not downloaded again."""
    if kind not in SECTION_KINDS:
        raise EmptySourceError("That section cannot be rewritten.")
    resource = store.get_resource(resource_id)
    if resource.status != STATUS_OK or not resource.extracted_text.strip():
        raise EmptySourceError("There is no source text to study.")
    active_settings = settings or get_settings()
    caller = complete or OpenAIModelClient(active_settings).complete
    request = StudyRequest(
        title=resource.title,
        url=resource.url,
        headings=[],
        source_text=resource.extracted_text,
        level=resource.level,
        goal=resource.goal,
        minutes=resource.minutes,
    )
    outline = _saved_outline(store, resource_id)
    if kind != "study" and not outline:
        outline = write_outline(request, caller)
    try:
        section = _one_section(
            kind,
            request,
            outline,
            caller,
            shortened=_was_shortened(store, resource_id),
        )
    except AgentError as exc:
        store.save_artifact(
            resource_id=resource_id,
            kind=kind,
            status=STATUS_ERROR,
            error_message=str(exc),
            markdown="",
            model_name=active_settings.openai_model,
        )
        if kind == "lesson":
            store.replace_lesson_checks(resource_id, [])
        return open_saved(store, resource_id)
    store.save_artifact(
        resource_id=resource_id,
        kind=kind,
        status=STATUS_OK,
        markdown=section.markdown,
        model_name=active_settings.openai_model,
    )
    if kind == "lesson":
        store.replace_lesson_checks(
            resource_id,
            [(item.question, item.answer, item.excerpt) for item in section.hidden_checks],
        )
    return open_saved(store, resource_id)


def _saved_outline(store: StudyStore, resource_id: int) -> str:
    study = next(
        (artifact for artifact in store.list_artifacts(resource_id) if artifact.kind == "study"),
        None,
    )
    if study is None:
        return ""
    return _outline_text(study.markdown)


def _was_shortened(store: StudyStore, resource_id: int) -> bool:
    study = next(
        (artifact for artifact in store.list_artifacts(resource_id) if artifact.kind == "study"),
        None,
    )
    return study is not None and study.markdown.startswith(SHORTENED_NOTICE)


def _outline_text(markdown: str) -> str:
    if markdown.startswith(SHORTENED_NOTICE):
        markdown = markdown[len(SHORTENED_NOTICE) :].lstrip("\n")
    marker = "## Outline\n\n"
    if not markdown.startswith(marker):
        return ""
    rest = markdown[len(marker) :]
    if "\n\n## " in rest:
        return rest.split("\n\n## ", 1)[0].strip()
    return rest.strip()


def _one_section(
    kind: str,
    request: StudyRequest,
    outline: str,
    complete: Completer,
    *,
    shortened: bool = False,
) -> AgentSection:
    if kind == "study":
        result = write_study(request, complete=complete, outline=outline or None)
        return _with_outline(
            AgentSection("study", STATUS_OK, "", result.markdown, list(result.checks)),
            result.outline,
            shortened=shortened,
        )
    if kind == "lesson":
        result = write_lesson(request, outline, complete)
        hidden = tuple(
            HiddenCheck(section.check_question, section.answer, section.quote)
            for section in result.sections
            if section.check_question.strip()
        )
        return AgentSection(
            "lesson",
            STATUS_OK,
            "",
            result.markdown,
            list(result.checks),
            tuple(item.answer for item in hidden),
            hidden,
        )
    if kind == "lab":
        result = write_lab(request, outline, complete)
    elif kind == "project":
        result = write_project(request, outline, complete)
    else:
        result = write_runbook(request, outline, complete)
    return AgentSection(kind, STATUS_OK, "", result.markdown, list(result.checks))
